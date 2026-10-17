import os
import shutil
import tempfile
import pytest
import openpyxl
import pandas as pd

from logic import extract_statement
from settings import Settings
from tests.excel_comparator import compare_excel_workbooks

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REKKOR_DIR = os.path.join(BASE_DIR, "Rekkor")
SCRATCH_DIR = "/Users/michaelarhdyn/.gemini/antigravity-ide/brain/fa8c9aa9-d258-43c7-a53d-3295ddcb8831/scratch"
if not os.path.exists(SCRATCH_DIR):
    SCRATCH_DIR = os.path.join(BASE_DIR, "scratch")


class TestSettingsRegressionAndEquality:
    """Verifikasi Fase 4: Regresi default 100%, variasi kombinasi pengaturan, dan kesetaraan."""

    def test_default_regression_livin_scraping(self):
        """Memverifikasi keluaran default Livin.pdf identik 100% dengan baseline Fase 0."""
        pdf_path = os.path.join(REKKOR_DIR, "Livin.pdf")
        baseline_excel = os.path.join(SCRATCH_DIR, "baseline_livin_scraping.xlsx")

        if not os.path.exists(pdf_path) or not os.path.exists(baseline_excel):
            pytest.skip("Livin.pdf atau baseline Fase 0 tidak ditemukan.")

        with tempfile.TemporaryDirectory() as tmpdir:
            out_excel = os.path.join(tmpdir, "livin_default_test.xlsx")
            count = extract_statement(
                pdf_path,
                out_excel,
                module_name="parser.Livin",
                settings=Settings()
            )
            assert count == 141

            result = compare_excel_workbooks(baseline_excel, out_excel)
            print(f"\n[Livin Scraping] Sel dibandingkan: {result.total_cells}, Perbedaan: {len(result.differences)}")
            assert result.is_identical, f"Perbedaan ditemukan: {result.differences}"
            assert len(result.differences) == 0

    def test_default_regression_mandiri_scraping(self):
        """Memverifikasi keluaran default Mandiri.pdf identik 100% dengan baseline Fase 0."""
        pdf_path = os.path.join(REKKOR_DIR, "Mandiri.pdf")
        baseline_excel = os.path.join(SCRATCH_DIR, "baseline_mandiri_scraping.xlsx")

        if not os.path.exists(pdf_path) or not os.path.exists(baseline_excel):
            pytest.skip("Mandiri.pdf atau baseline Fase 0 tidak ditemukan.")

        with tempfile.TemporaryDirectory() as tmpdir:
            out_excel = os.path.join(tmpdir, "mandiri_default_test.xlsx")
            count = extract_statement(
                pdf_path,
                out_excel,
                module_name="parser.Mandiri",
                settings=Settings()
            )
            assert count == 190

            result = compare_excel_workbooks(baseline_excel, out_excel)
            print(f"\n[Mandiri Scraping] Sel dibandingkan: {result.total_cells}, Perbedaan: {len(result.differences)}")
            assert result.is_identical, f"Perbedaan ditemukan: {result.differences}"
            assert len(result.differences) == 0

    def test_default_regression_livin_ocr(self):
        """Memverifikasi keluaran default Livin_ocr.pdf identik 100% dengan baseline Fase 0."""
        pdf_path = os.path.join(REKKOR_DIR, "Livin_ocr.pdf")
        baseline_excel = os.path.join(SCRATCH_DIR, "baseline_livin_ocr.xlsx")

        if not os.path.exists(pdf_path) or not os.path.exists(baseline_excel):
            pytest.skip("Livin_ocr.pdf atau baseline Fase 0 tidak ditemukan.")

        with tempfile.TemporaryDirectory() as tmpdir:
            out_excel = os.path.join(tmpdir, "livin_ocr_default_test.xlsx")
            count = extract_statement(
                pdf_path,
                out_excel,
                module_name="parser.Livin",
                settings=Settings()
            )
            assert count == 93

            result = compare_excel_workbooks(baseline_excel, out_excel)
            print(f"\n[Livin OCR] Sel dibandingkan: {result.total_cells}, Perbedaan: {len(result.differences)}")
            assert result.is_identical, f"Perbedaan ditemukan: {result.differences}"
            assert len(result.differences) == 0

    def test_five_non_default_combinations_structure_and_types(self):
        """Memverifikasi 5 kombinasi non-default: tipe sel, format nilai, dan integritas nilai numerik."""
        pdf_path = os.path.join(REKKOR_DIR, "Livin.pdf")
        if not os.path.exists(pdf_path):
            pytest.skip("Livin.pdf tidak ditemukan.")

        combinations = [
            # 1: Standar Indonesia (Ribuan titik, desimal koma, 2 digit, minus, teks)
            (
                "Kombinasi 1 (Format Indonesia)",
                Settings(
                    thousands_separator=".",
                    decimal_separator=",",
                    decimal_places=2,
                    negative_format="-1234",
                    output_number_type="text",
                    date_format="DD/MM/YYYY"
                ),
                "text"
            ),
            # 2: Ribuan Spasi (thousands ' ', decimal ',', 2 digit, minus, teks)
            (
                "Kombinasi 2 (Ribuan Spasi)",
                Settings(
                    thousands_separator=" ",
                    decimal_separator=",",
                    decimal_places=2,
                    negative_format="-1234",
                    output_number_type="text",
                    date_format="DD-MM-YYYY"
                ),
                "text"
            ),
            # 3: Desimal 0 digit (thousands '.', decimal ',', 0 digit, minus, teks)
            (
                "Kombinasi 3 (Desimal 0 Digit)",
                Settings(
                    thousands_separator=".",
                    decimal_separator=",",
                    decimal_places=0,
                    negative_format="-1234",
                    output_number_type="text",
                    date_format="DD MMM YYYY"
                ),
                "text"
            ),
            # 4: Negatif Tanda Kurung (thousands ',', decimal '.', 2 digit, kurung, teks)
            (
                "Kombinasi 4 (Negatif Kurung)",
                Settings(
                    thousands_separator=",",
                    decimal_separator=".",
                    decimal_places=2,
                    negative_format="(1234)",
                    output_number_type="text",
                    date_format="YYYY-MM-DD"
                ),
                "text"
            ),
            # 5: Mode Numerik Murni (thousands ',', decimal '.', 2 digit, numerik OpenXML)
            (
                "Kombinasi 5 (Mode Numerik Murni)",
                Settings(
                    thousands_separator=",",
                    decimal_separator=".",
                    decimal_places=2,
                    negative_format="-1234",
                    output_number_type="numeric",
                    date_format="YYYY-MM-DD"
                ),
                "numeric"
            ),
        ]

        with tempfile.TemporaryDirectory() as tmpdir:
            for name, s, expected_type in combinations:
                out_path = os.path.join(tmpdir, f"test_{s.output_number_type}_{s.decimal_places}.xlsx")
                count = extract_statement(
                    pdf_path,
                    out_path,
                    module_name="parser.Livin",
                    settings=s
                )
                assert count == 141, f"{name}: Jumlah baris tidak sesuai"

                wb = openpyxl.load_workbook(out_path, data_only=False)
                ws = wb.active

                # Ambil sel kolom Balance baris 2 (transaksi pertama)
                cell_balance = ws["F2"]
                cell_date = ws["A2"]

                if expected_type == "text":
                    assert cell_balance.data_type == "s", f"{name}: Data type harus string (s)"
                    val = str(cell_balance.value)
                    if s.thousands_separator == ".":
                        assert "." in val, f"{name}: Pemisah ribuan titik tidak ditemukan di '{val}'"
                    elif s.thousands_separator == " ":
                        assert " " in val, f"{name}: Pemisah ribuan spasi tidak ditemukan di '{val}'"
                    if s.decimal_places > 0:
                        assert s.decimal_separator in val, f"{name}: Pemisah desimal tidak ditemukan di '{val}'"
                    else:
                        assert s.decimal_separator not in val, f"{name}: Pemisah desimal muncul padahal 0 digit di '{val}'"

                    if s.date_format == "YYYY-MM-DD":
                        assert "-" in str(cell_date.value), f"{name}: Format tanggal ISO tidak sesuai di '{cell_date.value}'"
                else:
                    # expected_type == "numeric"
                    assert cell_balance.data_type == "n", f"{name}: Data type harus numeric (n)"
                    assert isinstance(cell_balance.value, (int, float)), f"{name}: Nilai harus int atau float"
                    assert cell_balance.number_format == "#,##0.00", f"{name}: Number format OpenXML harus '#,##0.00'"

                print(f"PASS: {name} terverifikasi valid (Tipe: {cell_balance.data_type}, Sampel: {cell_balance.value})")

    def test_negative_comparator_detection(self):
        """Negative test pembanding: Pastikan modifikasi satu sel terdeteksi secara presisi."""
        baseline_excel = os.path.join(SCRATCH_DIR, "baseline_livin_scraping.xlsx")
        if not os.path.exists(baseline_excel):
            pytest.skip("Baseline Livin scraping tidak ditemukan.")

        with tempfile.TemporaryDirectory() as tmpdir:
            tampered = os.path.join(tmpdir, "tampered.xlsx")
            shutil.copyfile(baseline_excel, tampered)

            wb = openpyxl.load_workbook(tampered)
            ws = wb.active
            ws["C5"].value = "MUTASI_MODIFIKASI_UJI"
            wb.save(tampered)

            res = compare_excel_workbooks(baseline_excel, tampered)
            assert not res.is_identical
            assert len(res.differences) >= 1
            assert any("C5" in d["location"] or "R5C3" in d["location"] for d in res.differences)
            print(f"\n[Negative Test] Berhasil mendeteksi {len(res.differences)} perbedaan pada lokasi C5/R5C3.")
