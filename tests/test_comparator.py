import os
import shutil
import tempfile
import openpyxl
import pytest
from tests.excel_comparator import compare_excel_workbooks

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REF_EXCEL = os.path.join(BASE_DIR, "scratch", "ref_Livin.xlsx")
# Fallback to Rekkor/Mandiri.xlsx if scratch ref not found
if not os.path.exists(REF_EXCEL):
    REF_EXCEL = os.path.join(BASE_DIR, "Rekkor", "Mandiri.xlsx")


def test_comparator_identical_files():
    if not os.path.exists(REF_EXCEL):
        pytest.skip("Reference Excel file not available for test.")

    res = compare_excel_workbooks(REF_EXCEL, REF_EXCEL)
    assert res.is_identical
    assert len(res.differences) == 0
    assert res.total_cells > 0


def test_comparator_negative_detects_value_diff():
    if not os.path.exists(REF_EXCEL):
        pytest.skip("Reference Excel file not available for test.")

    with tempfile.TemporaryDirectory() as tmpdir:
        altered_file = os.path.join(tmpdir, "altered.xlsx")
        shutil.copyfile(REF_EXCEL, altered_file)

        wb = openpyxl.load_workbook(altered_file)
        ws = wb.active
        # Alter a single cell value
        original_val = ws.cell(2, 2).value
        ws.cell(2, 2).value = "MODIFIED_VALUE_FOR_NEGATIVE_TEST"
        wb.save(altered_file)

        res = compare_excel_workbooks(REF_EXCEL, altered_file)
        assert not res.is_identical
        assert len(res.differences) >= 1
        diff_cats = [d["category"] for d in res.differences]
        assert "CELL_VALUE" in diff_cats
        diff_locations = [d["location"] for d in res.differences]
        assert any("R2C2" in loc for loc in diff_locations)
